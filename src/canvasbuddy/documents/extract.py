"""Turning course documents into text a model can read.

Dispatch is on **magic bytes, not the file extension**, because the extension is not
reliable here -- Canvas serves whatever the instructor uploaded, and content types are
frequently wrong.

Three formats, all of which appear in one real term:

* ``.docx`` -- a zip container. Read paragraphs *and tables*: the grading scheme is almost
  always a table, and a paragraph-only reader silently returns a syllabus with no weights.
* ``.pdf`` -- straightforward text extraction.
* ``.doc`` -- the pre-2007 binary format, an OLE compound file. There is no clean parser,
  so text runs are scraped heuristically out of the ``WordDocument`` stream. The output has
  formatting noise in it, which is acceptable because the consumer is a language model
  rather than a parser.
"""

from __future__ import annotations

import io
import logging
import re

log = logging.getLogger(__name__)

_ZIP_MAGIC = b"PK\x03\x04"
_PDF_MAGIC = b"%PDF"
_OLE_MAGIC = b"\xd0\xcf\x11\xe0"

#: Vocabulary that distinguishes a syllabus from everything else in a course's files.
#: Weighted because "term test" appearing at all is far more telling than "due".
_SIGNALS: dict[str, int] = {
    "midterm": 5,
    "mid-term": 5,
    "term test": 5,
    "final exam": 5,
    "grading scheme": 5,
    "marking scheme": 5,
    "weight": 3,
    "worth": 2,
    "quiz": 2,
    "assessment": 2,
    "participation": 2,
    "late policy": 3,
    "academic integrity": 2,
    "due": 1,
}
_PERCENT = re.compile(r"\d{1,3}\s?%")

#: Below this a document is stored but never sent to the model. The WileyPLUS
#: registration card that sits in MGAB03 named like a syllabus scores 0.
SYLLABUS_SCORE_FLOOR = 12

_WS = re.compile(r"[ \t]+")
_BLANKS = re.compile(r"\n{3,}")


class ExtractionError(RuntimeError):
    pass


def sniff(data: bytes) -> str:
    """Identify a document by its leading bytes."""
    if data.startswith(_ZIP_MAGIC):
        return "docx"
    if data.startswith(_PDF_MAGIC):
        return "pdf"
    if data.startswith(_OLE_MAGIC):
        return "doc"
    return "unknown"


def _tidy(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r", "\n")
    text = _WS.sub(" ", text)
    return _BLANKS.sub("\n\n", text).strip()


def _from_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    # Tables carry the grading scheme. Rendering rows pipe-separated keeps the
    # association between an assessment and its weight, which a flat dump loses.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return _tidy("\n".join(parts))


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return _tidy("\n".join((page.extract_text() or "") for page in reader.pages))


def _from_doc(data: bytes) -> str:
    """Best-effort scrape of a pre-2007 binary Word document.

    Word 97 interleaves text with formatting structures, so runs are pulled out by
    pattern: printable ASCII sequences, plus UTF-16LE sequences (which appear as
    alternating character/NUL bytes). Both are needed -- documents mix them.
    """
    import olefile

    if not olefile.isOleFile(io.BytesIO(data)):
        raise ExtractionError("not an OLE compound file")

    ole = olefile.OleFileIO(io.BytesIO(data))
    try:
        if not ole.exists("WordDocument"):
            raise ExtractionError("no WordDocument stream")
        stream = ole.openstream("WordDocument").read()
    finally:
        ole.close()

    ascii_runs = [m.decode("latin-1") for m in re.findall(rb"[\x20-\x7e]{6,}", stream)]
    wide_runs = [
        m.decode("utf-16-le", "ignore") for m in re.findall(rb"(?:[\x20-\x7e]\x00){6,}", stream)
    ]
    return _tidy("\n".join(ascii_runs + wide_runs))


_EXTRACTORS = {"docx": _from_docx, "pdf": _from_pdf, "doc": _from_doc}


def extract_text(data: bytes) -> tuple[str, str]:
    """Return ``(kind, text)`` for a document, raising if the format is unsupported."""
    kind = sniff(data)
    extractor = _EXTRACTORS.get(kind)
    if extractor is None:
        raise ExtractionError(f"unsupported format (leading bytes {data[:4]!r})")
    return kind, extractor(data)


def content_score(text: str) -> int:
    """How much this reads like a syllabus.

    Deliberately scores the *content*, never the filename. MGAB03 contains a file called
    "MGAB03 L04 - Kong - Fall 2026.pdf" which is a textbook registration card; trusting
    its name would hand the model a document with no assessment information in it and
    invite invented exam dates.
    """
    if not text:
        return 0
    low = text.lower()
    score = sum(weight for term, weight in _SIGNALS.items() if term in low)
    # A grading scheme means several percentages, not one stray figure.
    score += min(len(_PERCENT.findall(low)), 6)
    return score


def looks_like_syllabus(text: str) -> bool:
    return content_score(text) >= SYLLABUS_SCORE_FLOOR
